from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
import os

out = os.path.join(os.path.dirname(__file__), "IMS_API_Issues.xlsx")
desk = os.path.join(os.path.expanduser("~"), "Desktop", "IMS_API_Issues_Awai.xlsx")

wb = Workbook()
ws = wb.active
ws.title = "IMS Issues"
ws.sheet_view.rightToLeft = True

headers = [
    "#",
    "المنتج / Product",
    "الخطورة / Severity",
    "المشكلة / Issue",
    "التفاصيل / Details",
    "Endpoint",
    "ما نرسله / What we send",
    "سلوك IMS الحالي / Current IMS behavior",
    "المطلوب / Expected fix",
    "المسؤول / Owner",
    "الحالة / Status",
    "تاريخ الرصد / Observed",
    "أدلة / Evidence",
]

rows = [
    [
        1,
        "LISA – سيارات إجباري",
        "عالية / High",
        "طباعة الوثيقة من شاشة IMS لا تفتح PDF للوثائق الصادرة عبر API",
        "عند إصدار وثيقة من موقع Awai عبر API، زر طباعة داخل IMS يفتح صفحة IIS فارغة بدل ملف PDF. الملف موجود على in.eidc.gov.ly ويعمل من موقعنا ومن reprint API.",
        "POST /api/motor/lisa/create\nGET /api/motor/lisa/reprint/{code}",
        "Body إنشاء LISA حسب الوثائق + استلام pdfUrl",
        "يحفظ الوثيقة لكن لا يملأ مسار PDF لزر الطباعة في واجهة IMS → يفتح root IIS",
        "عند الإنشاء عبر API: حفظ نفس pdfUrl في صف الوثيقة وربطه بزر طباعة كما للوثائق اليدوية",
        "فريق IMS",
        "مفتوح / Open",
        "2026-09",
        "مثال: LISA-9F4V-5N37-89TA-PS4M\nPDF يعمل على in.eidc.gov.ly\nطباعة IMS → صفحة IIS فارغة",
    ],
    [
        2,
        "Travel – مسافرين",
        "حرجة / Critical",
        "إنشاء الوثيقة يرجع 400 للعميل بينما الوثيقة تُسجَّل وتعمل في IMS",
        "طلبات تظهر كفشل من جهة Awai (HTTP 400 Bad Request بدون JSON واضح) ثم يؤكد فريق IMS أنها مسجّلة وتعمل. الموقع يعرض فشلاً ولا يحفظ IssuedPolicy / لا يعرض PDF للمستخدم رغم نجاح الإصدار فعلياً.",
        "POST /api/travel/policy/create",
        "حقول الوثائق: travelerName, travelerLocation, travelerPhone, passportNumber, dateOfBirth, genderId, nationalityId, periodId, clauseId, countryId, destinationId, insuranceFrom",
        "أحياناً يرد 400 بنص Bad Request فقط (ليس {error:...}) ومع ذلك تُنشأ الوثيقة في IMS",
        "1) إرجاع 200 + بيانات الوثيقة عند النجاح الحقيقي\n2) عدم إنشاء الوثيقة عند الفشل\n3) رسالة خطأ JSON واضحة عند الرفض الحقيقي",
        "فريق IMS",
        "مفتوح / Open – مؤكد من الفريق 2026-10-08",
        "2026-10-07",
        "سجلات Awai: create returned 400. Response: Bad Request\nالفريق: نفس الطلبات مسجّلة وتعمل في IMS",
    ],
    [
        3,
        "Travel – مسافرين",
        "عالية / High",
        "تعارض الوثائق مع السلوك: destinationId دولة مقابل حزمة التغطية",
        "الوثائق تعرض countryId و destinationId مختلفين. الاختبار: countryId=1 (Schengen) + destinationId=56 (Germany) → 400. النجاح عندما destinationId = معرف حزمة areacover.",
        "POST /api/travel/policy/create\nPOST /api/travel/policy/draft\nPOST /api/travel/premium/enquiry",
        "اتفاق مؤقت: تجاهل الدولة؛ إرسال countryId=destinationId=حزمة التغطية (All Schengen)",
        "جدول الأسعار: الوجهة = All Schengen Area Countries. قائمة الدول لا تُسعَّر حالياً لـ create",
        "توضيح في الوثائق أو دعم دولة حقيقية في destinationId مع countryId للحزمة",
        "فريق IMS + Awai (مؤقتاً تجاهل الدولة)",
        "عمل مؤقت / Workaround",
        "2026-10-07",
        "premium dest=1 → 1026؛ dest=56 → لا قسط\ncreate dest=56 → 400؛ dest=1 → Success",
    ],
    [
        4,
        "Travel – مسافرين",
        "متوسطة / Medium",
        "رسائل خطأ غير مطابقة لشكل الوثائق",
        "الوثائق تقول أخطاء بشكل {\"error\":\"...\"}. الواقع غالباً نص خام Bad Request.",
        "Travel POST عند الرفض",
        "JSON + Bearer + Content-Type application/json",
        "400 بدون تفاصيل الحقل → صعوبة التشخيص",
        "إرجاع JSON موحد مع سبب الرفض",
        "فريق IMS",
        "مفتوح / Open",
        "2026-10-07",
        "Travelapidocs.html Error Response Format",
    ],
    [
        5,
        "Travel – مسافرين",
        "متوسطة / Medium",
        "لا قسط إذا لم تُضبط المدد والأسعار في لوحة IMS",
        "lookup/periods كانت فارغة ثم أُضيفت 30 يوماً. يلزم صف سعر مطابق (فترة + تغطية + حزمة + عمر).",
        "GET /api/travel/lookup/periods\nPOST /api/travel/premium/enquiry",
        "periodId / clauseId / destinationId=حزمة",
        "بدون سعر: hasPremium=false – No premium found for the selected combination",
        "الإبقاء على تهيئة الأسعار؛ توثيق التركيبات المدعومة",
        "فريق IMS (إعداد)",
        "محلول جزئياً – مدة 30 + سعر Schengen صحي 1026",
        "2026-10-07",
        "سعر نشط في اللوحة: شهر / 1-95 / All Schengen / تأمين صحي / 1026",
    ],
    [
        6,
        "Travel – مسافرين",
        "منخفضة / Low",
        "تسمية متضاربة: حزمة vs دولة ونفس id=1",
        "areacover id=1 = Schengen؛ destinations id=1 = Libya بالصدفة — يسبب لبس عند التشخيص.",
        "Lookups + لوحة الأسعار",
        "N/A",
        "تسميات مربكة بين الواجهة والـ API",
        "توحيد التسميات في الواجهة والوثائق",
        "فريق IMS",
        "مفتوح / Open",
        "2026-10-07",
        "areacover 1 = All Schengen\ndestinations 1 = Libya",
    ],
    [
        7,
        "عام / All APIs",
        "متوسطة / Medium",
        "Test يمنع Create/Finalize (403)",
        "سلوك موثّق: Test للمعاينة؛ الإصدار يحتاج Production key لكل منتج.",
        "POST .../create و .../finalize",
        "X-API-Key حسب المنتج",
        "403 في Test",
        "استخدام مفاتيح Production للإصدار",
        "Awai ops + IMS",
        "معلوم / Known",
        "2026-09",
        "Travelapidocs.html Environment",
    ],
]

header_fill = PatternFill("solid", fgColor="1B3A5F")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
severe = {
    "حرجة / Critical": PatternFill("solid", fgColor="F8D7DA"),
    "عالية / High": PatternFill("solid", fgColor="FFF3CD"),
    "متوسطة / Medium": PatternFill("solid", fgColor="CFE2FF"),
    "منخفضة / Low": PatternFill("solid", fgColor="E2E3E5"),
}
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
wrap = Alignment(wrap_text=True, vertical="top", readingOrder=2)

for col, h in enumerate(headers, 1):
    cell = ws.cell(1, col, h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center", readingOrder=2)
    cell.border = thin

for r_i, row in enumerate(rows, 2):
    for c_i, val in enumerate(row, 1):
        cell = ws.cell(r_i, c_i, val)
        cell.alignment = wrap
        cell.border = thin
        cell.font = Font(name="Calibri", size=10)
    sev = row[2]
    if sev in severe:
        ws.cell(r_i, 3).fill = severe[sev]

widths = [4, 18, 14, 28, 40, 28, 32, 36, 36, 14, 22, 12, 36]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30
for r in range(2, 2 + len(rows)):
    ws.row_dimensions[r].height = 95
ws.auto_filter.ref = f"A1:M{1 + len(rows)}"
ws.freeze_panes = "A2"

ws2 = wb.create_sheet("Summary")
ws2.sheet_view.rightToLeft = True
ws2["A1"] = "Awai ↔ IMS – سجل مشاكل التكامل"
ws2["A1"].font = Font(bold=True, size=14, color="1B3A5F")
ws2["A2"] = f"تاريخ الملف: {date.today().isoformat()}"
ws2["A3"] = "الغرض: مشاركة الملاحظات مع فريق IMS"
ws2["A5"] = "قرارات مؤقتة من جهة Awai"
ws2["A5"].font = Font(bold=True)
ws2["A6"] = "1) تجاهل حقل الدولة / traveler country destination مؤقتاً."
ws2["A7"] = "2) إرسال حزمة التغطية (areacover مثل All Schengen) في countryId و destinationId معاً."
ws2["A8"] = "3) اعتبار أن 400 من travel create قد يعني وثيقة صدرت فعلاً — يحتاج إصلاح IMS (#2)."
ws2["A10"] = "أولوية الإصلاح عند IMS"
ws2["A10"].font = Font(bold=True)
ws2["A11"] = "#2 Travel create يرد 400 رغم نجاح الإصدار (حرج)"
ws2["A12"] = "#1 طباعة LISA من IMS لا تفتح PDF للوثائق عبر API (عالٍ)"
ws2["A13"] = "#3/#6 توضيح destinationId حزمة vs دولة"
ws2["A14"] = "#4 رسائل خطأ JSON واضحة"
ws2.column_dimensions["A"].width = 110

ws3 = wb.create_sheet("Issues (EN)")
en_headers = [
    "#", "Product", "Severity", "Issue", "Details", "Endpoint",
    "What we send", "Current IMS behavior", "Expected fix", "Owner", "Status", "Observed", "Evidence",
]
en_rows = [
    [1, "LISA Compulsory Motor", "High",
     "IMS Print does not open PDF for API-issued policies",
     "API-created policies exist; PDF on in.eidc.gov.ly works from Awai. IMS Print opens empty IIS page because PDF path is not stored on the policy row for API creates.",
     "POST /api/motor/lisa/create", "LISA create body; receive pdfUrl",
     "Policy saved; print path empty", "Persist API pdfUrl and wire Print", "IMS", "Open", "2026-09",
     "LISA-9F4V-5N37-89TA-PS4M"],
    [2, "Travel", "Critical",
     "create returns HTTP 400 while policy is actually issued in IMS",
     "Awai shows failure and does not save local IssuedPolicy, but IMS confirms those requests are registered and working.",
     "POST /api/travel/policy/create", "Documented create JSON fields",
     "Plain 'Bad Request'; policy still created",
     "Return 200 on real success; no create on failure; clear JSON errors", "IMS",
     "Open – confirmed 2026-10-08", "2026-10-07", "App logs 400; IMS has policies"],
    [3, "Travel", "High",
     "Docs vs live: country destinationId rejected",
     "Docs show different countryId/destinationId. Live: only areacover package id works. Temporary: ignore traveler country field.",
     "create / draft / premium/enquiry", "countryId=destinationId=package (Schengen)",
     "Country id (e.g. Germany 56) rejected", "Document or support country destinationId", "IMS + Awai workaround",
     "Workaround", "2026-10-07", "dest=1 OK; dest=56 fail"],
    [4, "Travel", "Medium", "Error body not as documented",
     "Docs promise {error}; live often raw Bad Request", "Travel POST failures", "Valid JSON + Bearer",
     "Opaque 400", "Structured errors", "IMS", "Open", "2026-10-07", "Travelapidocs.html"],
    [5, "Travel", "Medium", "Empty periods / missing rates block premium",
     "Need periods + matching price row", "lookup/periods + premium/enquiry", "period/clause/package",
     "hasPremium false until configured", "Maintain rate catalog", "IMS config", "Partial", "2026-10-07",
     "30-day Schengen health 1026"],
    [6, "Travel", "Low", "Naming collision package vs country id=1",
     "areacover 1=Schengen; destinations 1=Libya", "Lookups + pricing UI", "N/A", "Confusing labels",
     "Align naming", "IMS", "Open", "2026-10-07", ""],
    [7, "All", "Medium", "Test key blocks finalize/create (403)",
     "Documented Test vs Production", "create/finalize", "Product API keys", "403 in Test",
     "Use Production keys to issue", "Ops", "Known", "2026-09", "Docs Environment"],
]
for col, h in enumerate(en_headers, 1):
    cell = ws3.cell(1, col, h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = thin
fills = {
    "Critical": PatternFill("solid", fgColor="F8D7DA"),
    "High": PatternFill("solid", fgColor="FFF3CD"),
    "Medium": PatternFill("solid", fgColor="CFE2FF"),
    "Low": PatternFill("solid", fgColor="E2E3E5"),
}
for r_i, row in enumerate(en_rows, 2):
    for c_i, val in enumerate(row, 1):
        cell = ws3.cell(r_i, c_i, val)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = thin
    if row[2] in fills:
        ws3.cell(r_i, 3).fill = fills[row[2]]
for i, w in enumerate([4, 18, 12, 28, 40, 26, 28, 32, 32, 12, 22, 12, 28], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.row_dimensions[1].height = 28
for r in range(2, 9):
    ws3.row_dimensions[r].height = 85
ws3.auto_filter.ref = f"A1:M{1 + len(en_rows)}"
ws3.freeze_panes = "A2"

wb.save(out)
wb.save(desk)
print(out)
print(desk)
